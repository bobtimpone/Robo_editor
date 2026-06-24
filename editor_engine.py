import docx
import urllib.request
import urllib.error
import json
import time

try:
    import tiktoken
except ImportError:
    raise ImportError("The 'tiktoken' library is missing. Run: pip install tiktoken")

class EditorEngine:
    def __init__(self, ollama_url="http://localhost:11434"):
        self.ollama_url = ollama_url
        self.cancel_requested = False
        
        try:
            self.tokenizer = tiktoken.get_encoding("cl100k_base")
        except Exception as e:
            print(f"Warning: Tokenizer initialization failed. {e}")
            self.tokenizer = None

    def _count_tokens(self, text):
        if self.tokenizer:
            return len(self.tokenizer.encode(text))
        return len(text.split()) * 1.3 

    def _para_to_markdown(self, para):
        md_text = ""
        for run in para.runs:
            text = run.text
            if not text:
                continue
            
            left_space = " " if text.startswith(" ") else ""
            right_space = " " if text.endswith(" ") else ""
            clean_text = text.strip()
            
            if not clean_text:
                md_text += text
                continue

            if run.bold and run.italic:
                md_text += f"{left_space}***{clean_text}***{right_space}"
            elif run.bold:
                md_text += f"{left_space}**{clean_text}**{right_space}"
            elif run.italic:
                md_text += f"{left_space}*{clean_text}*{right_space}"
            else:
                md_text += f"{left_space}{clean_text}{right_space}"
                
        return md_text

    def _markdown_to_runs(self, md_text, new_para):
        import re
        parts = re.split(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*)', md_text)
        
        for part in parts:
            if not part:
                continue
            
            if part.startswith('***') and part.endswith('***'):
                run = new_para.add_run(part[3:-3])
                run.bold = True
                run.italic = True
            elif part.startswith('**') and part.endswith('**'):
                run = new_para.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = new_para.add_run(part[1:-1])
                run.italic = True
            else:
                new_para.add_run(part)

    def load_and_chunk(self, filepath, max_tokens_per_chunk=1500):
        doc = docx.Document(filepath)
        chunks = []
        current_chunk = []
        current_token_count = 0
        
        for para in doc.paragraphs:
            if not para.text.strip():
                continue
                
            md_text = self._para_to_markdown(para)
            tokens = self._count_tokens(md_text)
            
            current_chunk.append(md_text)
            current_token_count += tokens
            
            if current_token_count >= max_tokens_per_chunk:
                chunks.append("\n\n".join(current_chunk))
                current_chunk = []
                current_token_count = 0
                
        if current_chunk:
            chunks.append("\n\n".join(current_chunk))
            
        return chunks

    def save_document(self, edited_chunks, output_filepath):
        new_doc = docx.Document()
        for chunk in edited_chunks:
            paragraphs = chunk.split('\n\n')
            for p in paragraphs:
                if p.strip():
                    new_para = new_doc.add_paragraph()
                    self._markdown_to_runs(p.strip(), new_para)
                    
        new_doc.save(output_filepath)
        return output_filepath

    def call_ollama(self, model, system_prompt, text, context_size=8192, temperature=0.3, stream_callback=None, max_retries=3, is_edit_pass=True):
        url = f"{self.ollama_url}/api/generate"
        
        if is_edit_pass:
            final_prompt = (
                f"{system_prompt}\n\n"
                "CRITICAL INSTRUCTION: You are receiving text formatted in Markdown. "
                "You MUST preserve any bold (**text**) or italic (*text*) formatting in your output. "
                "Return the revised text first. If you found issues you could not fix, "
                "append '---EXCEPTIONS---' on a new line, followed by a bulleted list of those issues.\n\n"
                "TEXT TO EDIT:\n"
                f"{text}"
            )
        else:
            final_prompt = f"{system_prompt}\n\nTEXT:\n{text}"
        
        data = {
            "model": model,
            "prompt": final_prompt,
            "stream": True,
            "options": {"temperature": temperature, "num_ctx": context_size}
        }
        
        payload = json.dumps(data).encode('utf-8')
        headers = {'Content-Type': 'application/json'}
        
        attempt = 0
        base_delay = 2 
        
        while attempt <= max_retries:
            try:
                req = urllib.request.Request(url, data=payload, headers=headers)
                generated_text = ""
                
                with urllib.request.urlopen(req, timeout=600) as response:
                    for line in response:
                        if self.cancel_requested:
                            raise Exception("Processing cancelled by user.")
                            
                        if line:
                            chunk_data = json.loads(line.decode('utf-8'))
                            new_words = chunk_data.get("response", "")
                            generated_text += new_words
                            
                            if stream_callback:
                                stream_callback(generated_text)
                                
                            if chunk_data.get("done"):
                                return generated_text.strip()
                                
            except (urllib.error.URLError, ConnectionResetError, TimeoutError) as e:
                attempt += 1
                if attempt > max_retries:
                    raise Exception(f"Ollama API failed after {max_retries} attempts. Final Error: {str(e)}")
                
                delay = base_delay ** attempt
                print(f"Network error: {str(e)}. Retrying in {delay} seconds... (Attempt {attempt}/{max_retries})")
                time.sleep(delay)
                
            except Exception as e:
                raise e